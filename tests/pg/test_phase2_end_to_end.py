"""Phase 2 acceptance, end to end on a throwaway local PostgreSQL with synthetic data only:

  workbook upload -> stored cells -> TL answers -> HRBP override -> newer upload un-flags a day
  -> effective status per day -> reconciled payroll export (Pending / Missing explicit).

Covers, in one period: an explicit Present, an explicit Absent, an HRBP override to a leave, a
leave in the source, a blank (missing) cell, an unresolved case, a review-required source change
and an unmapped CRM with no flagged day (D4)."""
import datetime
import hashlib
import importlib.util
import io
import json

import openpyxl
import pytest
from pathlib import Path
from psycopg.rows import dict_row

from app import data, ingest_flow
from app.attendance_view import hrbp_period_attendance, tl_period_attendance
from app.effective import BLANK, MISSING, PENDING, REVIEW_REQUIRED
from app.report import reconcile
from app.verdicts import VERDICT_LABEL
from ingestion.reference import parse_reference
from ingestion.summary import UNMAPPED_NO_FLAGS, ingest_summary

_spec = importlib.util.spec_from_file_location(        # the top-level tests/conftest.py builders
    "tests_conftest", Path(__file__).resolve().parents[1] / "conftest.py")
_top = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_top)
_build_reference = _top._build_reference

D = datetime.date
P = D(2026, 9, 15)
HEADER = ["CRM", "Normal Days", "Abnormal Days", "16-Sep", "17-Sep", "18-Sep", "19-Sep"]
FIRST = [["E-1", 1, 2, "Absent", "Absent", "Normal", None],
         ["E-2", 1, 3, "No Show", "Annual Leave", "Absent", "Absent"],
         ["E-GHOST", 4, 0, "Normal", "Normal", "Normal", "Normal"]]
# the newer upload: identical, except E-2's 19-Sep is no longer flagged
SECOND = [r if r[0] != "E-2" else r[:6] + ["Normal"] for r in FIRST]


def _workbook(tmp_path, name, rows):
    wb = openpyxl.Workbook()
    _build_reference(wb)
    sm = wb.create_sheet("Summary Report")
    sm.append(["Enhanced Attendance Summary Report"])
    sm.append([])
    sm.append(HEADER)
    for r in rows:
        sm.append(r)
    path = tmp_path / name
    wb.save(path)
    return path


def _upload(conn, path):
    ref = parse_reference(str(path))
    res = ingest_summary(str(path), ref, year=2026)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    conn.commit()
    summary = ingest_flow.load_confirmed(conn, ref, res, path.name, "hrbp:test", sha)
    conn.commit()
    return ref, res, summary


def _case_id(conn, crm, day):
    with conn.cursor() as cur:
        cur.execute("select c.id from attendance.cases c join attendance.employees e "
                    "on e.id = c.employee_id where e.crm = %s and c.work_date = %s", (crm, day))
        return cur.fetchone()[0]


def _manager(conn, crm):
    with conn.cursor(row_factory=dict_row) as cur:
        cur.execute("select id, crm from attendance.managers where crm = %s", (crm,))
        return cur.fetchone()


@pytest.fixture()
def flow(conn, tmp_path):
    first = _workbook(tmp_path, "Attendance_first.xlsx", FIRST)
    second = _workbook(tmp_path, "Attendance_second.xlsx", SECOND)
    ref, res, summary = _upload(conn, first)
    tl = "tl:TL-A"
    p1 = _case_id(conn, "E-1", D(2026, 9, 16))
    data.add_attachment(conn, p1, f"cases/{p1}/crm.png", "crm.png", "image/png", 10)
    assert data.submit_verdict(conn, p1, "present", None, "calls seen", tl)
    assert data.submit_verdict(conn, _case_id(conn, "E-2", D(2026, 9, 16)), "absent", None,
                               "no show confirmed", tl)
    assert data.close_case(conn, _case_id(conn, "E-1", D(2026, 9, 17)), "hrbp:test",
                           final_status="sick_leave", comment="medical certificate")
    _upload(conn, second)
    return {"first": first, "second": second, "res": res, "summary": summary, "ref": ref}


def test_upload_stores_every_cell_and_flags_the_right_days(conn, flow):
    res = flow["res"]
    assert len(res.cells) == 3 * 4                                   # employees x dates
    assert sorted((c.employee_crm, c.work_date.day) for c in res.cases) == \
        [("E-1", 16), ("E-1", 17), ("E-2", 16), ("E-2", 18), ("E-2", 19)]
    with conn.cursor() as cur:
        cur.execute("select count(*), bool_and(full_cells) from attendance.ingestion_runs")
        assert cur.fetchone() == (2, True)
        cur.execute("select count(*) from attendance.attendance_days")
        assert cur.fetchone()[0] == 2 * 3 * 4


def test_effective_status_for_every_kind_of_day(conn, flow):
    grid = {(r["employee_crm"], r["work_date"].day): r for r in hrbp_period_attendance(conn, P)}
    want = {("E-1", 16): ("Present", "tl"),
            ("E-1", 17): ("Sick Leave", "hrbp_override"),
            ("E-1", 18): ("Normal", "original"),
            ("E-1", 19): (BLANK, "original"),
            ("E-2", 16): ("Absent", "tl"),
            ("E-2", 17): ("Annual Leave", "original"),
            ("E-2", 18): (PENDING, "case"),
            ("E-2", 19): (REVIEW_REQUIRED, "case"),
            ("E-1", 20): (MISSING, "none")}                  # no upload covers it
    got = {k: (grid[k]["effective_status"], grid[k]["effective_source"]) for k in want}
    assert got == want
    assert [k for k, r in grid.items() if r["needs_review"]] == [("E-2", 19)]


def test_d4_unmapped_crm_is_reported_once_and_never_absent(conn, flow):
    assert flow["summary"].exceptions >= 1
    listed = data.list_unmapped_no_flags(conn)
    assert [r["crm"] for r in listed] == ["E-GHOST"]
    assert listed[0]["source_filename"] == "Attendance_second.xlsx"   # newest upload first
    with conn.cursor() as cur:
        cur.execute("select count(*) from attendance.ingestion_exceptions where reason = %s",
                    (UNMAPPED_NO_FLAGS,))
        assert cur.fetchone()[0] == 2                                  # one per upload
        cur.execute("select count(*) from attendance.cases c join attendance.employees e "
                    "on e.id = c.employee_id where e.crm = 'E-GHOST'")
        assert cur.fetchone()[0] == 0
    preview = ingest_flow.build_preview(flow["res"], flow["ref"])
    assert preview["unmapped_no_flags"] == 1


def test_reconciled_export_reflects_effective_status_and_marks_pending(conn, flow):
    closed, unresolved = data.list_closed_cases(conn), data.list_unresolved_cases(conn)
    result = reconcile(io.BytesIO(flow["second"].read_bytes()), closed, VERDICT_LABEL, year=2026,
                       staff_gaps=data.list_reference_gaps(conn), unresolved_cases=unresolved)
    wb = openpyxl.load_workbook(io.BytesIO(result.xlsx))
    ws = wb["Summary Report"]
    matrix = {row[0]: row[3:7] for row in ws.iter_rows(min_row=4, values_only=True) if row[0]}
    assert matrix["E-1"] == ("Present", "Sick Leave", "Normal", None)    # blank stays blank
    assert matrix["E-2"] == ("Absent", "Annual Leave", "Pending", "Normal")
    assert matrix["E-GHOST"] == ("Normal",) * 4
    if "Staff not in Attendance" in wb.sheetnames:                     # D4 stays off that sheet
        assert "E-GHOST" not in {r[0] for r in wb["Staff not in Attendance"].iter_rows(
            values_only=True)}
    assert (result.overwritten, result.pending_marked, result.pending_not_marked,
            result.blank_cells, result.not_in_workbook) == (3, 1, 1, 1, 0)

    changes = list(wb["Changes"].iter_rows(min_row=2, values_only=True))
    assert len(changes) == len(closed) + len(unresolved) == 5        # every case, none dropped
    by_key = {(r[0], r[2]): r for r in changes}
    assert by_key[("E-2", "2026-09-19")][4] == "Pending"
    assert by_key[("E-2", "2026-09-19")][9] == "No — unresolved; cell is not Absent"
    assert by_key[("E-1", "2026-09-17")][4] == "Sick Leave"

    # reconciliation: every flagged day in the file is accounted for exactly once
    flagged = sum(1 for r in FIRST for v in r[3:] if v in ("Absent", "No Show"))
    assert flagged == result.overwritten + result.unresolved


def test_export_after_the_period_close_is_unchanged_and_still_pending(conn, flow):
    # the close decides nothing, so the export built after it is identical, Pending included
    def build():
        r = reconcile(io.BytesIO(flow["second"].read_bytes()), data.list_closed_cases(conn),
                      VERDICT_LABEL, year=2026, staff_gaps=data.list_reference_gaps(conn),
                      unresolved_cases=data.list_unresolved_cases(conn))
        ws = openpyxl.load_workbook(io.BytesIO(r.xlsx))["Summary Report"]
        return r._replace(xlsx=b""), [row for row in ws.iter_rows(values_only=True)]
    before = build()
    res = data.close_open_month(conn, "hrbp:test")
    assert res["pending"] + res["kept_for_review"] == before[0].unresolved
    after = build()
    assert after == before
    assert after[0].unresolved == 2 and after[0].pending_marked == 1

def test_tl_view_stays_case_only_and_label_only(conn, flow):
    mgr = _manager(conn, "TL-A")
    rows = data.tl_cases_for_manager(conn, mgr)
    assert {r["status_label"] for r in rows} <= {"Absent", "No Show", "Flagged — review required"}
    dumped = json.dumps(rows, default=str)
    for secret in ("Sick", "certificate", "Annual", "hrbp:test", "medical"):
        assert secret not in dumped
    grid = json.dumps(tl_period_attendance(conn, mgr["id"], P), default=str)
    assert "medical certificate" not in grid and "effective_status" not in grid
    other = _manager(conn, "TL-C")
    assert data.tl_cases_for_manager(conn, other) == []                  # another team sees none
